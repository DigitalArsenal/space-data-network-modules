#!/usr/bin/env python3
"""Extract small NASA CARA numerical fixtures; no MATLAB runtime or network.

Run from the worktree root with .oracle/venv/bin/python. Only public numerical
vectors and derived encounter-plane projections are emitted; source code,
workbooks, PDFs and .mat data remain in .oracle. This implements the math of
Pc2D_Foster.m:92-113 (projection and eigenvalue repair), and independently checks
it with scipy double integration. Expected values come from NASA tests/tables.
"""
import argparse
import csv
import importlib.util
from pathlib import Path
import re
import subprocess
import sys
sys.dont_write_bytecode = True

import numpy as np
import openpyxl
import xlrd
from scipy.special import erfcinv

MODULE = Path(__file__).resolve().parents[1]
ROOT = MODULE.parents[1]
ORACLE = ROOT / '.oracle/CARA_Analysis_Tools'
FIXTURES = MODULE / 'src/cpp/tests/fixtures'
PROB = Path('DistributedMatlab/ProbabilityOfCollision')
FOSTER_TEST = PROB / 'UnitTests/Pc2D_Foster_UnitTest.m'
CIRCLE_TEST = PROB / 'UnitTests/PcCircle_UnitTest.m'
XCORR = Path('DistributedMatlab/Utils/CovarianceDeCorrelationCode')
FIELDS = ['case', 'xi_km', 'zeta_km', 'sxx_km2', 'sxz_km2', 'szz_km2',
          'radius_km', 'expected', 'relative_tolerance', 'absolute_tolerance',
          'repair_expected', 'source', 'epoch']
NUMBER = r'[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?'


def load_oracle():
    path = Path(__file__).with_name('pc-dblquad-oracle.py')
    spec = importlib.util.spec_from_file_location('dblquad_oracle', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def matlab_array(text, name):
    match = re.search(r'\b' + re.escape(name) + r'\s*=\s*\{?\s*\[(.*?)\]', text, re.S)
    if not match:
        raise ValueError(f'Missing MATLAB numeric array {name}')
    return np.array([float(n) for n in re.findall(NUMBER, match.group(1))])


def matlab_cell(text, name):
    match = re.search(r'\b' + name + r'\s*=\s*\{(.*?)\}', text, re.S)
    return [float(n) for n in re.findall(NUMBER, match.group(1))]


def line_of(text, token):
    return text[:text.index(token)].count('\n') + 1


def encounter_geometry(r1, v1, c1, r2, v2, c2, radius):
    """CARA Foster projection: input states km/km/s, covariance km²."""
    r, v = np.asarray(r1) - r2, np.asarray(v1) - v2
    y = v / np.linalg.norm(v)
    z = np.cross(r, v)
    z /= np.linalg.norm(z)
    x = np.cross(y, z)
    rotation = np.stack([x, z])
    c = rotation @ (c1[:3, :3] + c2[:3, :3]) @ rotation.T
    # CARA uses norm(r), including its sub-ms departure from exact TCA.
    return [np.linalg.norm(r), 0.0, c[0, 0], (c[0, 1] + c[1, 0]) / 2,
            c[1, 1], radius]


def read_alfano(number, oracle=ORACLE):
    p = oracle / 'DataFiles/AlfanoInputData/Alfano_2009_Test_Cases' / f'Case{number} data at TCA.xls'
    sheet = xlrd.open_workbook(p).sheet_by_index(0)
    vec = lambda row: np.array(sheet.row_values(row - 1)[:3], dtype=float) / 1000
    cov = lambda row: np.array([sheet.row_values(j)[:6] for j in range(row - 1, row + 5)], dtype=float) / 1e6
    return vec(5), vec(9), cov(13), vec(23), vec(27), cov(31)


def read_cdm(path):
    """Read only the numeric state/covariance and DCP fields needed by tests."""
    head, objects = {}, []
    active = head
    for line in path.read_text().splitlines():
        if '=' not in line:
            continue
        key, value = (part.strip() for part in line.split('=', 1))
        if key == 'OBJECT':
            active = {}
            objects.append(active)
        active[key] = value.split('[')[0].strip()
    return head, objects


def cdm_state_cov(obj):
    """Return km/km/s and km² blocks in one common inertial orientation.

    For ITRF, axes coincide with ITRF at TCA and velocity includes omega × r.
    A common inertial rotation cancels in Pc. Tiny EOP/precession terms are
    omitted only for the two upstream correlation tests (0.5% tolerance).
    """
    r = np.array([float(obj[k]) for k in ('X', 'Y', 'Z')])
    v = np.array([float(obj[k]) for k in ('X_DOT', 'Y_DOT', 'Z_DOT')])
    if obj['REF_FRAME'] == 'ITRF':
        v += np.cross([0, 0, 7.29211514670698e-5], r)
    elif obj['REF_FRAME'] != 'EME2000':
        raise ValueError(f'Unsupported fixture frame {obj["REF_FRAME"]}')
    radial = r / np.linalg.norm(r)
    normal = np.cross(r, v)
    normal /= np.linalg.norm(normal)
    basis = np.stack([radial, np.cross(normal, radial), normal], axis=1)
    labels = ['R', 'T', 'N', 'RDOT', 'TDOT', 'NDOT']
    c = np.zeros((6, 6))
    for i in range(6):
        for j in range(i + 1):
            key = f'C{labels[i]}_{labels[j]}'
            c[i, j] = c[j, i] = float(obj.get(key, 0)) / 1e6
    rotation = np.zeros((6, 6))
    rotation[:3, :3] = rotation[3:, 3:] = basis
    return r, v, rotation @ c @ rotation.T, basis


def fixture(name, geo, expected, rtol, atol, source, epoch, repair=0):
    return [name, *geo, expected, rtol, atol, repair, source, epoch]


def foster_cases(oracle):
    text = (oracle / FOSTER_TEST).read_text()
    data = [matlab_array(text, key) for key in ('r1', 'v1', 'cov1', 'r2', 'v2', 'cov2')]
    data[2], data[5] = data[2].reshape(3, 3), data[5].reshape(3, 3)
    geo = encounter_geometry(*data, .02)
    yield fixture('Foster-Omitron01', geo, 2.70601573490125e-5, .001, 0,
                  f'{FOSTER_TEST}:48-73', 'static J2000 state; epoch unspecified upstream')
    # The newer PcCircle test supplies a much tighter reference for the same case.
    yield fixture('PcCircle-Omitron01', geo, 2.706023476569787e-5, 1e-10, 0,
                  f'{CIRCLE_TEST}:46-66', 'static J2000 state; epoch unspecified upstream')
    for i, (hbr, pc, tol) in enumerate(zip(matlab_cell(text, 'AlfanoHBR'),
                                        matlab_cell(text, 'AlfanoExpSolution'),
                                        matlab_cell(text, 'AlfanoAccuracy')), 1):
        geo = encounter_geometry(*read_alfano(i, oracle), hbr / 1000)
        yield fixture(f'Foster-Alfano{i:02}', geo, pc, tol, 0,
                      f'{FOSTER_TEST}:31-43; Alfano_2009_Test_Cases/Case{i} data at TCA.xls rows 5 9 13:18 23 27 31:36',
                      'published at-TCA ECI state; time from epoch in workbook A1')
    p = Path('DataFiles/SampleCDMs/OmitronTestCase_Test07_NonPDCovariance.cdm')
    head, objects = read_cdm(oracle / p)
    a, b = [cdm_state_cov(obj) for obj in objects]
    geo = encounter_geometry(*a[:3], *b[:3], .020)
    yield fixture('Foster-NPD', geo, 0, 0, 1e-10, f'{FOSTER_TEST}:100-111; {p}:50-61;128-139',
                  head['TCA'] + ' UTC', repair=1)


def spreadsheet_cases(oracle):
    folder = Path('DataFiles/PcTestCaseCDMs')
    sheet = openpyxl.load_workbook(oracle / folder / 'CARA_PcMethod_Test_Conjunctions.xlsx',
                                  data_only=True, read_only=True).active
    # Both low- and high-speed states; a high-Pc case and known usage violations.
    for rownum in (2, 3, 14, 27, 49):
        row = next(sheet.iter_rows(min_row=rownum, max_row=rownum, values_only=True))
        head, objects = read_cdm(oracle / folder / (row[0] + '.cdm'))
        a, b = [cdm_state_cov(obj) for obj in objects]
        geo = encounter_geometry(*a[:3], *b[:3], row[4] / 1000)
        yield fixture(f'CARA-CDM-row{rownum}', geo, row[8], 1e-8, 0,
                      f'{folder}/CARA_PcMethod_Test_Conjunctions.xlsx!I{rownum}; {row[0]}.cdm',
                      head['TCA'] + ' UTC')


def correlation_cases(oracle):
    rows = []
    for number, expected in [(1, 8.16e-6), (3, 3.20e-5)]:
        path = XCORR / 'UnitTests/InputFiles' / f'DensityDecorrelationTestCaseCDM{number:02}.txt'
        head, objects = read_cdm(oracle / path)
        a, b = [cdm_state_cov(obj) for obj in objects]
        sensitivities = []
        for obj, item in zip(objects, (a, b)):
            sigma = float(obj['COMMENT DCP Density Forecast Uncertainty'])
            sensitivity = np.array([float(x) for x in obj['COMMENT DCP Sensitivity Vector RTN Pos'].split()]) / 1000
            sensitivities.append(sigma * (item[3] @ sensitivity))
        cross = np.outer(*sensitivities)
        corrected = a[2][:3, :3] + b[2][:3, :3] - cross - cross.T
        geo = encounter_geometry(a[0], a[1], corrected, b[0], b[1], np.zeros((3, 3)), .006)
        rows.append([f'CARA-XCorr{number:02}', *(a[0]-b[0]), *(a[1]-b[1]),
                     *(a[2]+b[2]).flatten(), *cross.flatten(), .006, expected, .005,
                     head['TCA'] + ' UTC', f'{path}; Pc_CovXCorr_UnitTest.m:{14 if number == 1 else 234}'])
        yield fixture(f'CARA-XCorr{number:02}', geo, expected, .005, 0,
                      f'{path}; Pc_CovXCorr_UnitTest.m:{14 if number == 1 else 234}',
                      head['TCA'] + ' UTC')
    fields = ['case'] + [f'r{i}_km' for i in range(3)] + [f'v{i}_km_s' for i in range(3)] + [f'c{i}{j}' for i in range(6) for j in range(6)] + [f'x{i}{j}' for i in range(3) for j in range(3)] + ['radius_km', 'expected', 'rtol', 'epoch', 'source']
    write_csv(FIXTURES / 'cara_correlation.csv', fields, rows)


def usage_cases(oracle):
    folder = Path('DataFiles/PcTestCaseCDMs')
    sheet = openpyxl.load_workbook(oracle / folder / 'CARA_PcMethod_Test_Conjunctions.xlsx',
                                  data_only=True, read_only=True).active
    rows = []
    for rownum in (2, 3, 14, 26, 27, 49, 52, 53):
        row = next(sheet.iter_rows(min_row=rownum, max_row=rownum, values_only=True))
        head, objects = read_cdm(oracle / folder / (row[0] + '.cdm'))
        a, b = [cdm_state_cov(obj) for obj in objects]
        rows.append([f'CARA-row{rownum}', *a[0], *a[1], *b[0], *b[1],
                     *a[2].flatten(), *b[2].flatten(), row[4]/1000, row[15],
                     row[10], row[11], row[12], row[13], row[14], head['TCA'] + ' UTC',
                     f'{folder}/CARA_PcMethod_Test_Conjunctions.xlsx row {rownum}; {row[0]}.cdm'])
    fields = ['case'] + [f'{v}{i}' for v in ('r1_', 'v1_', 'r2_', 'v2_') for i in range(3)] + [f'{c}{i}{j}' for c in ('c1_', 'c2_') for i in range(6) for j in range(6)] + ['radius_km', 'usage_code', 'nc2d', 'nc3d', 'sdmc', 'sdmc_lo', 'sdmc_hi', 'epoch', 'source']
    write_csv(FIXTURES / 'cara_usage.csv', fields, rows)


def hall_cases(oracle):
    test_path = PROB / 'UnitTests/Pc3D_Hall_UnitTest.m'
    text = (oracle / test_path).read_text()
    reference_path = PROB / 'Pc3D_Hall_Utils/Output_AlfanoTestCases_Pc3D/A09TestCases_01_12_Distributed.csv'
    with (oracle / reference_path).open() as f:
        published = list(csv.DictReader(f))
    rows = []
    for number, (radius, expected, tolerance, limit) in enumerate(zip(
            matlab_cell(text, 'AlfanoHBR'), matlab_cell(text, 'AlfanoExpSolution'),
            matlab_cell(text, 'AlfanoAccuracy'), matlab_cell(text, 'AlfanoTLimit')), 1):
        r1, v1, c1, r2, v2, c2 = read_alfano(number, oracle)
        reference = published[number - 1]
        rows.append([f'Hall-Alfano{number:02}', *r1, *v1, *r2, *v2,
                     *c1.flatten(), *c2.flatten(), radius / 1000, expected, tolerance,
                     -limit, limit, float(reference['CARA_PcMC']),
                     float(reference['CARA_PcMCLo95']), float(reference['CARA_PcMCHi95']),
                     'published at-TCA ECI state; time from epoch in workbook A1',
                     f'{test_path}:33-49; Case{number} data at TCA.xls rows 5 9 13:18 23 27 31:36; {reference_path}:{number+1}'])
    fields = ['case'] + [f'{v}{i}' for v in ('r1_', 'v1_', 'r2_', 'v2_') for i in range(3)] + [f'{c}{i}{j}' for c in ('c1_', 'c2_') for i in range(6) for j in range(6)] + ['radius_km', 'expected_hall', 'rtol', 'time_start_s', 'time_end_s', 'published_mc', 'published_mc_lo', 'published_mc_hi', 'epoch', 'source']
    write_csv(FIXTURES / 'cara_hall.csv', fields, rows)


def coppola_bounds(relative_r, relative_v, covariance, radius, gamma=1e-16):
    """Independent port of conj_bounds_Coppola.m:103-193; km and seconds."""
    speed = np.linalg.norm(relative_v)
    if speed < 100 * np.finfo(float).eps:
        raise ValueError('Undefined Coppola window at zero relative speed')
    x = relative_v / speed
    y = relative_r - x * np.dot(x, relative_r)
    y /= np.linalg.norm(y)
    rotation = np.stack([x, y, np.cross(x, y)])
    rotated_r = rotation @ relative_r
    a = rotation @ covariance[:3, :3] @ rotation.T
    w = a[1:3, 0]
    b = np.linalg.solve(a[1:3, 1:3], w)
    sigma = np.sqrt(max(0.0, 2 * (a[0, 0] - np.dot(b, w))))
    q0 = np.dot(b, rotated_r[1:3])
    b2 = np.dot(b, b)
    dmin, dmax = -radius * np.sqrt(b2), radius * np.sqrt(1 + b2)
    spread = erfcinv(gamma) * sigma
    return (-spread + q0 - dmax) / speed, (spread + q0 - dmin) / speed


def original_mc_cases(oracle):
    source = Path('DistributedMatlab/MonteCarloPc/UnitTests/Pc_MC_Kep2body_parallel_UnitTest.m')
    text = (oracle / source).read_text()
    matches = list(re.finditer(r'^        function test(\d+)\(testCase\)', text, re.M))
    rows = []
    for index, match in enumerate(matches):
        number = int(match.group(1))
        if number not in (1, 2, 3, 4, 5, 6, 10, 11):
            continue
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        block = text[match.start():end]
        active = '\n'.join(line.split('%', 1)[0] for line in block.splitlines())
        scalar = lambda name: float(re.search(r'^\s*' + name + r'\s*=\s*(' + NUMBER + ')', active, re.M).group(1))
        radius, expected, tolerance = scalar('HBR') / 1000, scalar('expSolution'), scalar('Accuracy')
        path = Path(re.search(r"(?:FileName|CDMfile)\s*=\s*'../../../([^']+)'", active).group(1))
        cdm_path = oracle / path
        if not cdm_path.exists():
            # Upstream test04 spells CrossTrack differently from the shipped
            # Crosstrack filename. Read exactly that file, never skip a case.
            alternatives = [p for p in cdm_path.parent.iterdir() if p.name.lower() == cdm_path.name.lower()]
            if len(alternatives) != 1:
                raise FileNotFoundError(cdm_path)
            cdm_path = alternatives[0]
            path = cdm_path.relative_to(oracle)
        head, objects = read_cdm(cdm_path)
        a, b = [cdm_state_cov(obj) for obj in objects]
        lower, upper = coppola_bounds(b[0] - a[0], b[1] - a[1], a[2] + b[2], radius)
        mode = 'Coppola gamma=1e-16'
        # Follow executable upstream code: test06 explicitly comments out
        # duration clipping. Tests10/11 replace the result with +/-1420 s.
        if re.search(r'^\s*if \(tau1-tau0\) > LowVelThreshold\*OrbitalPeriod', active, re.M):
            mu = scalar('GM') / 1e9
            semi_major = -mu / (np.dot(a[1], a[1]) - 2 * mu / np.linalg.norm(a[0]))
            period = 2 * np.pi * np.sqrt(semi_major ** 3 / mu)
            threshold = scalar('LowVelThreshold') * period
            if upper - lower > threshold:
                upper = threshold * upper / (upper - lower)
                lower = upper - threshold
                mode += '; clipped to 5% primary period'
        if number in (10, 11):
            lower, upper = -scalar('tau1'), scalar('tau1')
            mode = 'explicit +/-1420 s override'
        first = text[:match.start()].count('\n') + 1
        expected_match = re.search(r'^\s*expSolution\s*=', block, re.M)
        expected_line = first + block[:expected_match.start()].count('\n')
        cdm_text = cdm_path.read_text()
        state_lines = [str(i) for i, line in enumerate(cdm_text.splitlines(), 1)
                       if re.match(r'^X\s*=', line) or re.match(r'^CNDOT_NDOT\s*=', line)]
        provenance = f'{source}:{first}-{expected_line}; {path} state/cov blocks bounded by lines {" ".join(state_lines)}; {mode}'
        rows.append([f'CARA-MC-unit{number:02}', *a[0], *a[1], *b[0], *b[1],
                     *a[2].flatten(), *b[2].flatten(), radius, expected, tolerance,
                     lower, upper, head['TCA'] + ' UTC', provenance])
        print(f'CARA MC test{number:02}: expected={expected:.12g} rtol={tolerance:g} HBR_km={radius:g} window=[{lower:.12g};{upper:.12g}] {mode}')
    if len(rows) != 8:
        raise RuntimeError('Expected all 8 original Monte Carlo tests')
    fields = ['case'] + [f'{v}{i}' for v in ('r1_', 'v1_', 'r2_', 'v2_') for i in range(3)] + [f'{c}{i}{j}' for c in ('c1_', 'c2_') for i in range(6) for j in range(6)] + ['radius_km', 'expected_mc', 'rtol', 'time_start_s', 'time_end_s', 'epoch', 'source']
    write_csv(FIXTURES / 'cara_monte_carlo_unit.csv', fields, rows)


def write_csv(path, fields, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(fields)
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--oracle', type=Path, default=ORACLE)
    parser.add_argument('--skip-dblquad', action='store_true')
    args = parser.parse_args()
    rows = list(foster_cases(args.oracle)) + list(spreadsheet_cases(args.oracle)) + list(correlation_cases(args.oracle))
    oracle = load_oracle()
    maximum = 0.0
    print('CARA independent Python reference: case expected reference relative_error tolerance')
    for row in rows:
        _, x, z, xx, xz, zz, radius, expected, rtol, atol, repair, *_ = row
        value, _ = oracle.disk_probability([x, z], np.array([[xx, xz], [xz, zz]]), radius, repair=True)
        error = abs(value - expected)
        relative = error / expected if expected else 0.0
        maximum = max(maximum, relative)
        print(f'{row[0]:24s} {expected:.12e} {value:.12e} {relative:.6e} {rtol or atol:.1e}')
        if error > max(atol, rtol * abs(expected)):
            raise RuntimeError(f'Independent reference fails {row[0]}')
    print(f'CARA cases={len(rows)} max_relative_error={maximum:.9e}')
    write_csv(FIXTURES / 'cara_pc.csv', FIELDS, rows)
    usage_cases(args.oracle)
    hall_cases(args.oracle)
    original_mc_cases(args.oracle)
    if not args.skip_dblquad:
        subprocess.run([sys.executable, str(Path(__file__).with_name('pc-dblquad-oracle.py'))], check=True)
    frisbee = Path(__file__).with_name('regen-cara-frisbee-fixtures.py')
    if frisbee.exists():
        subprocess.run([sys.executable, str(frisbee), "--cara-root", str(args.oracle)], check=True)


if __name__ == '__main__':
    main()
